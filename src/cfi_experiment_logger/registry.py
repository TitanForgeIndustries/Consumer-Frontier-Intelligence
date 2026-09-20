"""Registry and spreadsheet/document exports."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .core import load_json


REGISTRY_FIELDS = [
    "experiment_id",
    "status",
    "started_at",
    "completed_at",
    "model",
    "method",
    "dataset",
    "baseline",
    "hypothesis",
    "change",
    "max_steps",
    "context_length",
    "final_step",
    "final_loss",
    "total_tokens",
    "tokens_per_second",
    "duration_seconds",
    "peak_vram_gb",
    "peak_system_ram_gb",
    "avg_gpu_utilization_pct",
    "peak_gpu_power_w",
    "peak_gpu_temperature_c",
    "training_event_count",
    "hardware_sample_count",
    "evaluation",
    "notes",
]


def _experiment_dirs(root: Path) -> list[Path]:
    if not root.exists():
        return []
    return sorted(
        path
        for path in root.iterdir()
        if path.is_dir() and (path / "summary.json").exists()
    )


def build_registry(root: str | Path = "experiments") -> list[dict[str, Any]]:
    root_path = Path(root)
    rows: list[dict[str, Any]] = []
    for path in _experiment_dirs(root_path):
        row = load_json(path / "summary.json")
        evaluation = row.get("evaluation") or {}
        row["evaluation"] = "; ".join(
            f"{key}={value}" for key, value in sorted(evaluation.items())
        )
        rows.append({field: row.get(field) for field in REGISTRY_FIELDS})
    return rows


def export_registry_csv(
    root: str | Path = "experiments",
    output: str | Path = "results/CFI_Experiment_Registry.csv",
) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows = build_registry(root)
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REGISTRY_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    return output_path


def export_registry_xlsx(
    root: str | Path = "experiments",
    output: str | Path = "results/CFI_Experiment_Registry.xlsx",
) -> Path:
    """Create a formatted workbook when openpyxl is installed."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "Excel export requires the optional 'excel' dependency. "
            "Install with: pip install -e '.[excel]'"
        ) from exc

    rows = build_registry(root)
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    experiments = workbook.active
    experiments.title = "Experiments"
    experiments.append(REGISTRY_FIELDS)
    for row in rows:
        experiments.append([row.get(field) for field in REGISTRY_FIELDS])

    _style_sheet(experiments, len(REGISTRY_FIELDS), len(rows) + 1, Table, TableStyleInfo, Font, PatternFill, get_column_letter)

    hardware = workbook.create_sheet("Hardware")
    hardware_headers = [
        "experiment_id",
        "peak_vram_gb",
        "peak_system_ram_gb",
        "avg_gpu_utilization_pct",
        "peak_gpu_power_w",
        "peak_gpu_temperature_c",
        "hardware_sample_count",
    ]
    hardware.append(hardware_headers)
    for row in rows:
        hardware.append([row.get(field) for field in hardware_headers])
    _style_sheet(hardware, len(hardware_headers), len(rows) + 1, Table, TableStyleInfo, Font, PatternFill, get_column_letter)

    training = workbook.create_sheet("Training")
    training_headers = [
        "experiment_id",
        "model",
        "method",
        "dataset",
        "max_steps",
        "context_length",
        "final_step",
        "final_loss",
        "total_tokens",
        "tokens_per_second",
        "duration_seconds",
    ]
    training.append(training_headers)
    for row in rows:
        training.append([row.get(field) for field in training_headers])
    _style_sheet(training, len(training_headers), len(rows) + 1, Table, TableStyleInfo, Font, PatternFill, get_column_letter)

    evaluation = workbook.create_sheet("Evaluation")
    evaluation.append(["experiment_id", "evaluation"])
    for row in rows:
        evaluation.append([row.get("experiment_id"), row.get("evaluation")])
    _style_sheet(evaluation, 2, len(rows) + 1, Table, TableStyleInfo, Font, PatternFill, get_column_letter)

    workbook.save(output_path)
    return output_path


def _style_sheet(sheet, width: int, height: int, Table, TableStyleInfo, Font, PatternFill, get_column_letter) -> None:
    header_fill = PatternFill("solid", fgColor="172033")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(width)}{height}"
    if height >= 2:
        table = Table(
            displayName=f"{sheet.title.replace(' ', '')}Table",
            ref=f"A1:{get_column_letter(width)}{height}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    for column_index in range(1, width + 1):
        values = [
            str(sheet.cell(row=row_index, column=column_index).value or "")
            for row_index in range(1, height + 1)
        ]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            42,
            max(12, max(len(value) for value in values) + 2),
        )
